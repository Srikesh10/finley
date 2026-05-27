"""
scripts/compare_models.py — Build a side-by-side model comparison report.

Reads all reports/queries_judged_*.xlsx files and produces:
  reports/model_comparison.xlsx

Run after all model benchmarks are complete.
"""

import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import glob
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

JUDGED_PATTERN = "reports/queries_judged_*.xlsx"
OUTPUT_FILE    = "reports/model_comparison.xlsx"

# ── Load all judged result files ──────────────────────────────────────────────

files = sorted(glob.glob(JUDGED_PATTERN))
if not files:
    print(f"No files found matching {JUDGED_PATTERN}")
    print("Run judge_queries.py for each model first.")
    sys.exit(1)

print(f"Found {len(files)} model result files:")
for f in files:
    print(f"  {f}")

# Parse each file into {model_label: [{n, question, verdict, score, notes}]}
model_data = {}

for filepath in files:
    label = os.path.basename(filepath).replace("queries_judged_", "").replace(".xlsx", "").upper()
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row - 1, values_only=True):
        if row[0] and isinstance(row[0], int):
            rows.append({
                "n":        row[0],
                "question": row[1],
                "answer":   row[2],
                "runtime":  row[3],
                "cost":     row[6],
                "verdict":  row[7],
                "score":    row[8],
                "notes":    row[9],
            })
    model_data[label] = rows
    passes    = sum(1 for r in rows if r["verdict"] == "PASS")
    avg_score = round(sum(r["score"] or 0 for r in rows) / len(rows), 1) if rows else 0
    print(f"  {label}: {passes}/{len(rows)} PASS | avg {avg_score}/10")

models = list(model_data.keys())
n_questions = len(next(iter(model_data.values())))

# ── Styles ────────────────────────────────────────────────────────────────────

THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

HDR_FILL  = fill("111827")
HDR_FONT  = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
PASS_FILL = fill("D1FAE5")
FAIL_FILL = fill("FEE2E2")
STRIPE    = fill("F9FAFB")
WHITE     = fill("FFFFFF")
SUM_FILL  = fill("E5E7EB")

CENTER = Alignment(horizontal="center", vertical="center")
LEFT   = Alignment(horizontal="left",   vertical="top", wrap_text=True)
RIGHT  = Alignment(horizontal="right",  vertical="center")

# ── Build workbook ────────────────────────────────────────────────────────────

wb = openpyxl.Workbook()

# ── Sheet 1: Summary ──────────────────────────────────────────────────────────

ws_sum = wb.active
ws_sum.title = "Summary"

ws_sum.merge_cells("A1:F1")
ws_sum["A1"].value = "Finley — Model Comparison Summary"
ws_sum["A1"].font  = Font(name="Calibri", bold=True, size=14, color="111827")
ws_sum["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws_sum.row_dimensions[1].height = 30

sum_headers = ["Model", "Pass Rate", "Avg Score", "Avg Runtime (s)", "Cost/Query ($)", "Total Cost ($)"]
for ci, h in enumerate(sum_headers, 1):
    c = ws_sum.cell(row=2, column=ci, value=h)
    c.fill      = HDR_FILL
    c.font      = HDR_FONT
    c.border    = BORDER
    c.alignment = CENTER
ws_sum.row_dimensions[2].height = 20

for ri, model in enumerate(models, start=3):
    rows = model_data[model]
    passes    = sum(1 for r in rows if r["verdict"] == "PASS")
    avg_score = round(sum(r["score"] or 0 for r in rows) / len(rows), 1) if rows else 0
    avg_rt    = round(sum(r["runtime"] or 0 for r in rows) / len(rows), 1) if rows else 0
    total_cost = round(sum(r["cost"] or 0 for r in rows), 4)
    cost_per_q = round(total_cost / len(rows), 5) if rows else 0

    base = STRIPE if ri % 2 == 0 else WHITE
    vals = [model, f"{passes}/{len(rows)} ({round(passes/len(rows)*100)}%)", avg_score, avg_rt, cost_per_q, total_cost]
    for ci, val in enumerate(vals, 1):
        c = ws_sum.cell(row=ri, column=ci, value=val)
        c.fill      = base
        c.border    = BORDER
        c.alignment = CENTER if ci > 1 else Alignment(horizontal="left", vertical="center")
        c.font      = Font(name="Calibri", size=10, bold=(ci == 1))

ws_sum.column_dimensions["A"].width = 14
for col in ["B", "C", "D", "E", "F"]:
    ws_sum.column_dimensions[col].width = 18

# ── Sheet 2: Side-by-side by question ────────────────────────────────────────

ws_detail = wb.create_sheet("Question by Question")

# Header row: #, Question, then per-model Score + Verdict pairs
headers = ["#", "Question"]
for m in models:
    headers += [f"{m} Score", f"{m} Verdict"]
headers.append("Winner")

for ci, h in enumerate(headers, 1):
    c = ws_detail.cell(row=1, column=ci, value=h)
    c.fill      = HDR_FILL
    c.font      = HDR_FONT
    c.border    = BORDER
    c.alignment = CENTER if ci != 2 else Alignment(horizontal="left", vertical="center")
ws_detail.row_dimensions[1].height = 20

for qi in range(n_questions):
    ri = qi + 2
    base = STRIPE if ri % 2 == 0 else WHITE
    question = next(iter(model_data.values()))[qi]["question"]
    n        = next(iter(model_data.values()))[qi]["n"]

    ws_detail.cell(row=ri, column=1, value=n).fill = base
    ws_detail.cell(row=ri, column=1).alignment = CENTER
    ws_detail.cell(row=ri, column=1).border    = BORDER

    c = ws_detail.cell(row=ri, column=2, value=question)
    c.fill      = base
    c.alignment = LEFT
    c.border    = BORDER

    scores = {}
    col = 3
    for m in models:
        row_data = model_data[m][qi]
        score   = row_data["score"]
        verdict = row_data["verdict"]
        scores[m] = score

        sc = ws_detail.cell(row=ri, column=col, value=score)
        sc.fill      = (PASS_FILL if verdict == "PASS" else FAIL_FILL) if verdict in ("PASS", "FAIL") else base
        sc.alignment = CENTER
        sc.border    = BORDER
        sc.font      = Font(name="Calibri", bold=True, size=10)

        vc = ws_detail.cell(row=ri, column=col + 1, value=verdict)
        vc.fill      = (PASS_FILL if verdict == "PASS" else FAIL_FILL) if verdict in ("PASS", "FAIL") else base
        vc.alignment = CENTER
        vc.border    = BORDER
        vc.font      = Font(name="Calibri", bold=True, size=10,
                             color="065F46" if verdict == "PASS" else "991B1B")
        col += 2

    winner = max(scores, key=scores.get) if scores else ""
    wc = ws_detail.cell(row=ri, column=col, value=winner)
    wc.fill      = base
    wc.alignment = CENTER
    wc.border    = BORDER
    ws_detail.row_dimensions[ri].height = 40

# Column widths
ws_detail.column_dimensions["A"].width = 5
ws_detail.column_dimensions["B"].width = 52
col_letter_idx = 3
for m in models:
    ws_detail.column_dimensions[get_column_letter(col_letter_idx)].width     = 12
    ws_detail.column_dimensions[get_column_letter(col_letter_idx + 1)].width = 12
    col_letter_idx += 2
ws_detail.column_dimensions[get_column_letter(col_letter_idx)].width = 14

ws_detail.freeze_panes = "A2"

# ── Sheet 3: Failures only ────────────────────────────────────────────────────

ws_fail = wb.create_sheet("FAILs Only")

fail_headers = ["#", "Question"] + [f"{m} Score / Verdict" for m in models]
for ci, h in enumerate(fail_headers, 1):
    c = ws_fail.cell(row=1, column=ci, value=h)
    c.fill      = HDR_FILL
    c.font      = HDR_FONT
    c.border    = BORDER
    c.alignment = CENTER if ci != 2 else Alignment(horizontal="left", vertical="center")

fail_ri = 2
for qi in range(n_questions):
    any_fail = any(model_data[m][qi]["verdict"] == "FAIL" for m in models)
    if not any_fail:
        continue
    base     = STRIPE if fail_ri % 2 == 0 else WHITE
    question = next(iter(model_data.values()))[qi]["question"]
    n        = next(iter(model_data.values()))[qi]["n"]

    ws_fail.cell(row=fail_ri, column=1, value=n).fill = base
    ws_fail.cell(row=fail_ri, column=1).alignment = CENTER
    ws_fail.cell(row=fail_ri, column=1).border    = BORDER

    c = ws_fail.cell(row=fail_ri, column=2, value=question)
    c.fill = base; c.alignment = LEFT; c.border = BORDER

    for ci, m in enumerate(models, start=3):
        r       = model_data[m][qi]
        verdict = r["verdict"]
        cell_val = f"{r['score']}/10 {verdict}"
        c = ws_fail.cell(row=fail_ri, column=ci, value=cell_val)
        c.fill      = FAIL_FILL if verdict == "FAIL" else PASS_FILL
        c.alignment = CENTER
        c.border    = BORDER
        c.font      = Font(name="Calibri", bold=True, size=10,
                           color="991B1B" if verdict == "FAIL" else "065F46")
    ws_fail.row_dimensions[fail_ri].height = 50
    fail_ri += 1

ws_fail.column_dimensions["A"].width = 5
ws_fail.column_dimensions["B"].width = 52
for ci in range(3, 3 + len(models)):
    ws_fail.column_dimensions[get_column_letter(ci)].width = 20

wb.save(OUTPUT_FILE)
print(f"\nComparison saved to {OUTPUT_FILE}")
print(f"Sheets: Summary | Question by Question | FAILs Only")
