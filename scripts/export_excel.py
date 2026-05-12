import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

with open("reports/cost_benchmark.json") as f:
    results = json.load(f)

wb = openpyxl.Workbook()

# ── Styles ────────────────────────────────────────────────────────────────────
DARK_HEADER = PatternFill("solid", fgColor="1F2937")
GREEN_FILL  = PatternFill("solid", fgColor="D1FAE5")
RED_FILL    = PatternFill("solid", fgColor="FEE2E2")
YELLOW_FILL = PatternFill("solid", fgColor="FEF9C3")
GREY_FILL   = PatternFill("solid", fgColor="F3F4F6")
WHITE_FONT  = Font(color="FFFFFF", bold=True)
BOLD        = Font(bold=True)
WRAP        = Alignment(wrap_text=True, vertical="top")
TOP         = Alignment(vertical="top")

def hdr(ws, row, col, value):
    c = ws.cell(row=row, column=col, value=value)
    c.font = WHITE_FONT
    c.fill = DARK_HEADER
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return c

def thin_border():
    s = Side(style="thin", color="D1D5DB")
    return Border(left=s, right=s, top=s, bottom=s)

# ── Sheet 1: Summary ──────────────────────────────────────────────────────────
ws1 = wb.active
ws1.title = "Summary"

ws1.merge_cells("A1:D1")
c = ws1["A1"]
c.value = "COST + QUALITY BENCHMARK — Initial Code vs Current Code"
c.font = Font(bold=True, size=14, color="FFFFFF")
c.fill = DARK_HEADER
c.alignment = Alignment(horizontal="center", vertical="center")
ws1.row_dimensions[1].height = 30

headers = ["Metric", "Initial Code\n(Sonnet + raw CSV)", "Current Code\n(Haiku + JSON + Opus judge)", "Delta"]
for ci, h in enumerate(headers, 1):
    hdr(ws1, 2, ci, h)
ws1.row_dimensions[2].height = 35

ini_scores  = [r["initial_judgment"].get("overall", 0) for r in results]
cur_scores  = [r["current_judgment"].get("overall", 0) for r in results]
ini_costs   = [r["initial"]["cost_usd"] for r in results]
cur_gen     = [r["current"]["cost_usd"] for r in results]
cur_judge   = [r["current_judgment"].get("_cost_usd", 0) for r in results]
cur_costs   = [g + j for g, j in zip(cur_gen, cur_judge)]
ini_passes  = sum(1 for r in results if r["initial_judgment"].get("verdict") == "PASS")
cur_passes  = sum(1 for r in results if r["current_judgment"].get("verdict") == "PASS")

avg_ini = round(sum(ini_scores) / len(ini_scores), 2)
avg_cur = round(sum(cur_scores) / len(cur_scores), 2)
tot_ini = sum(ini_costs)
tot_cur = sum(cur_costs)
ratio   = round(tot_ini / max(tot_cur, 0.0001), 1)

rows = [
    ("Avg Quality Score (0–10)", avg_ini, avg_cur, f"{avg_cur - avg_ini:+.2f}"),
    ("Pass Rate", f"{ini_passes}/{len(results)}", f"{cur_passes}/{len(results)}", f"{cur_passes - ini_passes:+d}"),
    ("Total Cost (5 queries)", f"${tot_ini:.4f}", f"${tot_cur:.4f}", f"{ratio}x cheaper"),
    ("Haiku Generation Cost", "N/A", f"${sum(cur_gen):.4f}", ""),
    ("Opus Judge Cost", "N/A", f"${sum(cur_judge):.4f}", ""),
    ("Avg Input Tokens / Query", f"{round(sum(r['initial']['input_tokens'] for r in results)/len(results)):,}", f"{round(sum(r['current']['input_tokens'] for r in results)/len(results)):,}", ""),
    ("CSV / Data Size Passed to Model", "430,691 chars (raw CSV)", "5,987 chars (JSON summary)", "71.9x smaller"),
]

for ri, (metric, ini, cur, delta) in enumerate(rows, 3):
    fill = GREY_FILL if ri % 2 == 0 else PatternFill()
    for ci, val in enumerate([metric, ini, cur, delta], 1):
        cell = ws1.cell(row=ri, column=ci, value=val)
        cell.fill = fill
        cell.alignment = TOP
        cell.border = thin_border()
        if ci == 1:
            cell.font = BOLD

ws1.column_dimensions["A"].width = 36
ws1.column_dimensions["B"].width = 22
ws1.column_dimensions["C"].width = 28
ws1.column_dimensions["D"].width = 18

# ── Sheet 2: Per-Question Scores ──────────────────────────────────────────────
ws2 = wb.create_sheet("Per-Question Scores")

ws2.merge_cells("A1:G1")
c = ws2["A1"]
c.value = "Quality Score Comparison — Per Question"
c.font = Font(bold=True, size=13, color="FFFFFF")
c.fill = DARK_HEADER
c.alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[1].height = 28

hdrs2 = ["#", "Question", "Initial Score", "Current Score", "Delta", "Initial Verdict", "Current Verdict"]
for ci, h in enumerate(hdrs2, 1):
    hdr(ws2, 2, ci, h)
ws2.row_dimensions[2].height = 28

for ri, r in enumerate(results, 3):
    ini_s = r["initial_judgment"].get("overall", 0)
    cur_s = r["current_judgment"].get("overall", 0)
    delta = cur_s - ini_s
    ini_v = r["initial_judgment"].get("verdict", "?")
    cur_v = r["current_judgment"].get("verdict", "?")

    row_data = [ri - 2, r["question"], ini_s, cur_s, f"{delta:+}", ini_v, cur_v]
    fill = GREEN_FILL if delta > 0 else (RED_FILL if delta < 0 else GREY_FILL)

    for ci, val in enumerate(row_data, 1):
        cell = ws2.cell(row=ri, column=ci, value=val)
        cell.border = thin_border()
        cell.alignment = TOP
        if ci in (3, 4, 5):
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="top")
        if ci in (6, 7):
            cell.fill = RED_FILL if val == "FAIL" else GREEN_FILL
            cell.alignment = Alignment(horizontal="center", vertical="top")

ws2.column_dimensions["A"].width = 5
ws2.column_dimensions["B"].width = 50
ws2.column_dimensions["C"].width = 15
ws2.column_dimensions["D"].width = 15
ws2.column_dimensions["E"].width = 10
ws2.column_dimensions["F"].width = 16
ws2.column_dimensions["G"].width = 16

# ── Sheet 3: Per-Question Cost ────────────────────────────────────────────────
ws3 = wb.create_sheet("Per-Question Cost")

ws3.merge_cells("A1:F1")
c = ws3["A1"]
c.value = "Cost Comparison — Per Question (full pipeline)"
c.font = Font(bold=True, size=13, color="FFFFFF")
c.fill = DARK_HEADER
c.alignment = Alignment(horizontal="center", vertical="center")
ws3.row_dimensions[1].height = 28

hdrs3 = ["#", "Question", "Initial Cost (Sonnet)", "Haiku Gen Cost", "Opus Judge Cost", "Total Current Cost", "X Cheaper"]
for ci, h in enumerate(hdrs3, 1):
    hdr(ws3, 2, ci, h)
ws3.row_dimensions[2].height = 28

for ri, r in enumerate(results, 3):
    ini_c  = r["initial"]["cost_usd"]
    gen_c  = r["current"]["cost_usd"]
    jdg_c  = r["current_judgment"].get("_cost_usd", 0)
    tot_c  = gen_c + jdg_c
    x      = round(ini_c / max(tot_c, 0.0001), 1)
    fill   = GREY_FILL if ri % 2 == 0 else PatternFill()

    row_data = [ri - 2, r["question"], f"${ini_c:.4f}", f"${gen_c:.4f}", f"${jdg_c:.4f}", f"${tot_c:.4f}", f"{x}x"]
    for ci, val in enumerate(row_data, 1):
        cell = ws3.cell(row=ri, column=ci, value=val)
        cell.fill = fill
        cell.border = thin_border()
        cell.alignment = TOP
        if ci == 7:
            cell.fill = GREEN_FILL
            cell.alignment = Alignment(horizontal="center", vertical="top")

ws3.column_dimensions["A"].width = 5
ws3.column_dimensions["B"].width = 50
ws3.column_dimensions["C"].width = 22
ws3.column_dimensions["D"].width = 18
ws3.column_dimensions["E"].width = 18
ws3.column_dimensions["F"].width = 20
ws3.column_dimensions["G"].width = 12

# ── Sheet 4: Answer Comparison ────────────────────────────────────────────────
ws4 = wb.create_sheet("Answer Comparison")

ws4.merge_cells("A1:C1")
c = ws4["A1"]
c.value = "Side-by-Side Answer Comparison"
c.font = Font(bold=True, size=13, color="FFFFFF")
c.fill = DARK_HEADER
c.alignment = Alignment(horizontal="center", vertical="center")
ws4.row_dimensions[1].height = 28

hdrs4 = ["Question", "Initial Answer (Sonnet + raw CSV)", "Current Answer (Haiku + JSON summary)"]
for ci, h in enumerate(hdrs4, 1):
    hdr(ws4, 2, ci, h)
ws4.row_dimensions[2].height = 28

for ri, r in enumerate(results, 3):
    ini_ans = r["initial"]["answer"][:1000]
    cur_ans = r["current"]["answer"][:1000]
    row_data = [r["question"], ini_ans, cur_ans]
    ws4.row_dimensions[ri].height = 120
    for ci, val in enumerate(row_data, 1):
        cell = ws4.cell(row=ri, column=ci, value=val)
        cell.alignment = WRAP
        cell.border = thin_border()
        if ci == 1:
            cell.font = BOLD
            cell.fill = GREY_FILL

ws4.column_dimensions["A"].width = 35
ws4.column_dimensions["B"].width = 60
ws4.column_dimensions["C"].width = 60

wb.save("reports/benchmark_comparison.xlsx")
print("Saved to reports/benchmark_comparison.xlsx")
