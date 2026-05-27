"""
scripts/run_hybrid.py — Benchmark with full hybrid routing.

Python Lookup  → python_answer() directly (zero LLM cost)
LLM data query → Haiku  ($1/$5 per MTok)
LLM advice     → Sonnet ($3/$15 per MTok)

Output: reports/queries_results_hybrid.xlsx
"""

import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import json
import time

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from anthropic import AnthropicBedrock
from dotenv import load_dotenv
load_dotenv()

from finley import process_transactions, SYSTEM_PROMPT, AWS_KEY, AWS_SECRET_KEY, AWS_REGION
from finley.router import classify, classify_complexity, python_answer
from config import TX_FILE, ACCOUNTS_FILE, RECURRING_FILE, SALARY_FILE
from scripts.utils import (
    DARK, SUM_CLR, BORDER, HDR_FONT, BODY_FONT, NUM_FONT,
    WRAP_TOP, CENTER_TOP, RIGHT_TOP, row_fill, write_title, write_header_row,
)

MODELS = {
    "haiku": {
        "id":       "us.anthropic.claude-haiku-4-5-20251001-v1:0",
        "in_rate":  1.00,
        "out_rate": 5.00,
    },
    "sonnet": {
        "id":       "us.anthropic.claude-sonnet-4-6",
        "in_rate":  3.00,
        "out_rate": 15.00,
    },
}

QUERIES_FILE = "queries_new.txt"
OUTPUT_FILE  = "reports/queries_results_hybrid.xlsx"
CHECKPOINT   = "reports/.run_hybrid_checkpoint.json"
MAX_RETRIES  = 3
RETRY_DELAY  = 5

PYTHON_CLR = PatternFill("solid", fgColor="DBEAFE")   # light blue
DATA_CLR   = PatternFill("solid", fgColor="D1FAE5")   # light green
ADVICE_CLR = PatternFill("solid", fgColor="FEF3C7")   # light amber

with open(QUERIES_FILE, encoding="utf-8") as f:
    raw = [line.strip() for line in f]
queries = [line for line in raw if line and line.lower() != "question"]
print(f"Loaded {len(queries)} queries from {QUERIES_FILE}")

results = []
if os.path.exists(CHECKPOINT):
    with open(CHECKPOINT, encoding="utf-8") as f:
        results = json.load(f)
    print(f"Resuming from checkpoint — {len(results)}/{len(queries)} already done")
done_indices = {r["n"] for r in results}

print("Processing transactions...", end="", flush=True)
t0 = time.time()
summary = process_transactions(
    tx_path        = TX_FILE,
    accounts_path  = ACCOUNTS_FILE  if os.path.exists(ACCOUNTS_FILE)  else None,
    recurring_path = RECURRING_FILE if os.path.exists(RECURRING_FILE) else None,
    salary_path    = SALARY_FILE    if os.path.exists(SALARY_FILE)    else None,
)
summary_json = json.dumps(summary)
print(f" done ({time.time() - t0:.1f}s) — {len(summary_json):,} chars")

client = AnthropicBedrock(
    aws_access_key=AWS_KEY,
    aws_secret_key=AWS_SECRET_KEY,
    aws_region=AWS_REGION,
)

os.makedirs("reports", exist_ok=True)
n_python = 0
n_haiku  = 0
n_sonnet = 0

for i, question in enumerate(queries, start=1):
    if i in done_indices:
        r = next(r for r in results if r["n"] == i)
        route = r.get("route", "llm")
        if route == "python":
            n_python += 1
        elif route == "haiku":
            n_haiku += 1
        else:
            n_sonnet += 1
        continue

    route = classify(question)

    if route == "python":
        t_start   = time.time()
        py_answer = python_answer(question, summary)
        elapsed   = round(time.time() - t_start, 4)
        if py_answer is not None:
            answer, inp, out, cost = py_answer, 0, 0, 0.0
            n_python += 1
            print(f"[{i:02d}/{len(queries)}] [PYTHON] {question[:55]}... {elapsed}s | $0.00000")
        else:
            route = "llm"

    if route == "llm":
        complexity = classify_complexity(question)
        route      = "haiku" if complexity == "data" else "sonnet"
        model_cfg  = MODELS[route]

        answer = inp = out = elapsed = None
        for attempt in range(1, MAX_RETRIES + 1):
            try:
                t_start  = time.time()
                response = client.messages.create(
                    model=model_cfg["id"],
                    max_tokens=1500,
                    system=SYSTEM_PROMPT,
                    messages=[{"role": "user", "content": f"Transaction summary:\n{summary_json}\n\nQuestion: {question}"}],
                )
                elapsed = round(time.time() - t_start, 2)
                inp     = response.usage.input_tokens
                out     = response.usage.output_tokens
                answer  = response.content[0].text.strip()
                break
            except Exception as e:
                if attempt == MAX_RETRIES:
                    print(f" FAILED: {e}")
                    answer, inp, out, elapsed = f"ERROR: {e}", 0, 0, 0.0
                else:
                    print(f" retry {attempt}...", end="", flush=True)
                    time.sleep(RETRY_DELAY * attempt)

        cost = round(
            (inp / 1_000_000 * model_cfg["in_rate"]) +
            (out / 1_000_000 * model_cfg["out_rate"]),
            5,
        )
        if route == "haiku":
            n_haiku += 1
        else:
            n_sonnet += 1
        tag = "HAIKU " if route == "haiku" else "SONNET"
        print(f"[{i:02d}/{len(queries)}] [{tag}] {question[:50]}... {elapsed}s | {inp}in/{out}out | ${cost:.5f}")

    results.append({
        "n": i, "question": question, "answer": answer, "route": route,
        "runtime_s": elapsed, "input_tokens": inp, "output_tokens": out, "cost_usd": cost,
    })
    with open(CHECKPOINT, "w", encoding="utf-8") as f:
        json.dump(results, f)

results.sort(key=lambda r: r["n"])

# ── Build Excel ───────────────────────────────────────────────────────────────

HEADERS    = ["#", "Question", "Answer", "Route", "Runtime (s)", "Input Tokens", "Output Tokens", "Cost ($)"]
COL_WIDTHS = [5, 42, 65, 9, 13, 14, 15, 11]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Results"

write_title(ws, f"Finley Hybrid Routing Results — {len(results)} questions — Python + Haiku + Sonnet", "A1:H1")
write_header_row(ws, HEADERS, left_cols={2})

def _row_color(route):
    if route == "python": return PYTHON_CLR
    if route == "haiku":  return DATA_CLR
    if route == "sonnet": return ADVICE_CLR
    return row_fill(0)

for ri, r in enumerate(results, start=3):
    base_fill = _row_color(r.get("route", ""))
    vals = [r["n"], r["question"], r["answer"], r.get("route", "").upper(),
            r["runtime_s"], r["input_tokens"], r["output_tokens"], round(r["cost_usd"], 5)]
    for ci, val in enumerate(vals, 1):
        c = ws.cell(row=ri, column=ci, value=val)
        c.fill   = base_fill
        c.border = BORDER
        if ci in (1, 4, 5, 6, 7, 8):
            c.font      = NUM_FONT
            c.alignment = RIGHT_TOP if ci == 8 else CENTER_TOP
        else:
            c.font      = BODY_FONT
            c.alignment = WRAP_TOP
    ws.row_dimensions[ri].height = 90

for ci, w in enumerate(COL_WIDTHS, 1):
    ws.column_dimensions[get_column_letter(ci)].width = w

sr = len(results) + 3
total_cost  = round(sum(r["cost_usd"] for r in results), 4)
avg_runtime = round(sum(r["runtime_s"] for r in results) / len(results), 2)
total_in    = sum(r["input_tokens"] for r in results)
total_out   = sum(r["output_tokens"] for r in results)

ws.cell(row=sr, column=1, value="TOTAL").font = NUM_FONT
ws.cell(row=sr, column=1).alignment = CENTER_TOP
ws.cell(row=sr, column=1).fill = SUM_CLR

route_summary = f"P:{n_python} H:{n_haiku} S:{n_sonnet}"
for ci, val in zip([4, 5, 6, 7, 8],
                   [route_summary, f"{avg_runtime}s avg", total_in, total_out, total_cost]):
    c = ws.cell(row=sr, column=ci, value=val)
    c.font      = NUM_FONT
    c.alignment = RIGHT_TOP if ci == 8 else CENTER_TOP
    c.fill      = SUM_CLR
    c.border    = BORDER

ws.freeze_panes = "A3"
wb.save(OUTPUT_FILE)

if os.path.exists(CHECKPOINT):
    os.remove(CHECKPOINT)

haiku_cost  = round(sum(r["cost_usd"] for r in results if r.get("route") == "haiku"),  4)
sonnet_cost = round(sum(r["cost_usd"] for r in results if r.get("route") == "sonnet"), 4)

print(f"\nDone.")
print(f"  Python : {n_python}/{len(results)} ({round(n_python/len(results)*100)}%) — $0.0000")
print(f"  Haiku  : {n_haiku}/{len(results)}  ({round(n_haiku/len(results)*100)}%) — ${haiku_cost:.4f}")
print(f"  Sonnet : {n_sonnet}/{len(results)}  ({round(n_sonnet/len(results)*100)}%) — ${sonnet_cost:.4f}")
print(f"  Total cost       : ${total_cost:.4f}")
print(f"  Cost per query   : ${round(total_cost/len(results), 5):.5f}")
print(f"  Saved to         : {OUTPUT_FILE}")
