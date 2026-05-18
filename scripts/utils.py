"""Shared Excel styling utilities for all scripts."""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


# ── Fills ─────────────────────────────────────────────────────────────────────
DARK     = PatternFill("solid", fgColor="1F2937")
STRIPE   = PatternFill("solid", fgColor="F9FAFB")
WHITE    = PatternFill("solid", fgColor="FFFFFF")
PASS_CLR = PatternFill("solid", fgColor="D1FAE5")
FAIL_CLR = PatternFill("solid", fgColor="FEE2E2")
SUM_CLR  = PatternFill("solid", fgColor="E5E7EB")

# ── Fonts ─────────────────────────────────────────────────────────────────────
HDR_FONT  = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="Calibri", size=10)
NUM_FONT  = Font(name="Calibri", size=10, color="374151")

# ── Borders ───────────────────────────────────────────────────────────────────
def thin_border(color: str = "E5E7EB") -> Border:
    s = Side(style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)

BORDER = thin_border()

# ── Alignments ────────────────────────────────────────────────────────────────
WRAP_TOP   = Alignment(wrap_text=True, vertical="top")
CENTER_TOP = Alignment(horizontal="center", vertical="top")
RIGHT_TOP  = Alignment(horizontal="right", vertical="top")
LEFT_MID   = Alignment(horizontal="left", vertical="center")


def row_fill(row_idx: int) -> PatternFill:
    return STRIPE if row_idx % 2 == 0 else WHITE


def write_header_row(ws, headers: list[str], row: int = 2,
                     center_cols: set = None, left_cols: set = None):
    """Write a styled dark header row."""
    center_cols = center_cols or set()
    left_cols   = left_cols or set()
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=ci, value=h)
        c.fill   = DARK
        c.font   = HDR_FONT
        c.border = BORDER
        c.alignment = LEFT_MID if ci in left_cols else CENTER_TOP
    ws.row_dimensions[row].height = 20


def write_title(ws, text: str, span: str, row: int = 1):
    """Merge cells and write a bold title row."""
    ws.merge_cells(span)
    c = ws[span.split(":")[0]]
    c.value     = text
    c.font      = Font(name="Calibri", bold=True, size=12, color="111827")
    c.alignment = LEFT_MID
    ws.row_dimensions[row].height = 26
